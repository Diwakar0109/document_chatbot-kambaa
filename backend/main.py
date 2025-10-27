import os
import io
import re
from dotenv import load_dotenv
from typing import List, Tuple, Dict

from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

# --- Document Parsers ---
from pypdf import PdfReader
from docx import Document
import pandas as pd

# --- LangChain Imports ---
from langchain_groq import ChatGroq
from langchain_cohere import CohereEmbeddings
from langchain.text_splitter import CharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain.chains import ConversationalRetrievalChain
from langchain.prompts import PromptTemplate ### CORRECTION 1: Added PromptTemplate import ###
from langchain_core.messages import SystemMessage, HumanMessage


load_dotenv()

# --- App Initialization ---
app = FastAPI(
    title="AI Chatbot Backend",
    description="Backend for a RAG chatbot with session management.",
    version="1.9.0" # Version bump for logic correction
)

# --- CORRECTED CORS CONFIGURATION ---
allowed_origins = [
    "https://asha-med.vercel.app",
    "http://127.0.0.1:5500",
    "http://localhost:5500",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# --- Global Variables & Constants ---
COHERE_EMBED_MODEL = "embed-english-light-v3.0"
embeddings = None
sessions: Dict[str, Dict] = {}


@app.on_event("startup")
async def startup_event():
    """
    Initializes the Cohere embeddings model on application startup.
    """
    print("--- Application Startup ---")
    global embeddings
    
    print(f"Initializing Cohere embeddings model ('{COHERE_EMBED_MODEL}')...")
    try:
        cohere_api_key = os.getenv("COHERE_API_KEY")
        if not cohere_api_key:
            raise ValueError("FATAL: COHERE_API_KEY environment variable not set.")
        
        embeddings = CohereEmbeddings(cohere_api_key=cohere_api_key, model=COHERE_EMBED_MODEL)
        _ = embeddings.embed_query("Test query")

        print(f"✅ Cohere embeddings model ('{COHERE_EMBED_MODEL}') initialized successfully.")
    except Exception as e:
        print(f"❌ FATAL ERROR: Could not initialize Cohere embeddings model: {e}")
        raise RuntimeError("Application startup failed: Embeddings could not be loaded.") from e
        
    print("--- Application Ready ---")


# --- Pydantic Models ---
class ChatRequest(BaseModel):
    session_id: str
    question: str
    chat_history: List[Tuple[str, str]] = []

class ChatResponse(BaseModel):
    answer: str

class SessionRequest(BaseModel):
    session_id: str


# --- Helper Functions ---
def extract_text_from_file(file: UploadFile, file_contents: bytes) -> str:
    """Extracts text from uploaded file based on its extension."""
    text = ""
    try:
        if file.filename.endswith(".pdf"):
            pdf_stream = io.BytesIO(file_contents)
            reader = PdfReader(pdf_stream)
            for page in reader.pages:
                text += page.extract_text() or ""
        elif file.filename.endswith(".docx"):
            docx_stream = io.BytesIO(file_contents)
            doc = Document(docx_stream)
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif file.filename.endswith(".txt"):
            text = file_contents.decode("utf-8")
        elif file.filename.endswith((".xlsx", ".xls")):
            excel_stream = io.BytesIO(file_contents)
            xls = pd.ExcelFile(excel_stream)
            all_sheets_text = []
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
                sheet_text = '\n'.join([' '.join(map(str, row)) for row in df.values])
                all_sheets_text.append(f"Sheet: {sheet_name}\n{sheet_text}")
            text = "\n\n---\n\n".join(all_sheets_text)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Please upload a .txt, .pdf, .docx, or .xlsx file.")
        
        if not text.strip():
            raise HTTPException(status_code=400, detail=f"Could not extract any text from '{file.filename}'. The file might be empty or scanned as an image.")
            
        return text
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error parsing file '{file.filename}': {e}")


# --- API Endpoints ---
@app.post("/session/clear", summary="Clear a session's chat history and knowledge base")
async def clear_session(request: SessionRequest):
    if request.session_id in sessions:
        sessions.pop(request.session_id, None)
        print(f"Cleared session: {request.session_id}")
    return {"status": "success", "message": "Session cleared."}


@app.post("/upload", summary="Upload a knowledge base file for a specific session")
async def upload_knowledge_base(session_id: str = Form(...), file: UploadFile = File(...)):
    if not session_id:
        raise HTTPException(status_code=400, detail="Session ID is missing.")
    if not embeddings:
        raise HTTPException(status_code=503, detail="Embeddings model is not available. Check server logs for startup errors.")

    if session_id not in sessions:
        sessions[session_id] = {"vector_store": None, "rag_chain": None, "uploaded_files": []}
    
    session = sessions[session_id]

    if file.filename in session["uploaded_files"]:
        raise HTTPException(status_code=400, detail=f"File '{file.filename}' has already been uploaded in this session.")

    try:
        file_contents = await file.read()
        text = extract_text_from_file(file, file_contents)
        
        text_splitter = CharacterTextSplitter(separator="\n", chunk_size=1000, chunk_overlap=200, length_function=len)
        chunks = text_splitter.split_text(text)
        if not chunks:
            raise HTTPException(status_code=400, detail="Could not create any text chunks from the document.")
        
        if session["vector_store"]:
            new_docs_vector_store = FAISS.from_texts(texts=chunks, embedding=embeddings)
            session["vector_store"].merge_from(new_docs_vector_store)
        else:
            session["vector_store"] = FAISS.from_texts(texts=chunks, embedding=embeddings)

        llm = ChatGroq(temperature=0.2, model_name="llama-3.1-8b-instant", groq_api_key=os.getenv("GROQ_API_KEY"))

        ### CORRECTION 2: Added a prompt to condense the question using chat history ###
        condense_question_prompt_template = """
        Given the following conversation and a follow up question, rephrase the follow up question to be a standalone question, in its original language.

        Chat History:
        {chat_history}
        Follow Up Input: {question}
        Standalone question:"""
        CONDENSE_QUESTION_PROMPT = PromptTemplate.from_template(condense_question_prompt_template)

        ### CORRECTION 3: Included {context} and {question} placeholders in the main prompt ###
        prompt_template = """
YOU ARE "DR. Asha" — AN EXPERT VIRTUAL HEALTH ASSISTANT. YOUR GOAL IS TO ANALYZE THE PROVIDED MEDICAL REPORT CONTEXT AND ANSWER THE USER'S QUESTION IN A COMPASSIONATE, CLEAR, AND RESPONSIBLE MANNER.

Follow these instructions precisely:
1.  **Analyze the Context**: Base your answer *only* on the information found in the 'Context from uploaded documents' section below. Do not use any outside knowledge.
2.  **Address the User's Question**: Directly answer the user's specific 'Question'.
3.  **Maintain Persona**: Speak like a calm, empathetic, and professional doctor. Acknowledge the user's concerns.
4.  **Do Not Diagnose**: Never give a definitive diagnosis. Instead, explain what the results mean and suggest consulting a doctor for a formal diagnosis and treatment plan.
5.  **Handle Missing Information**: If the context does not contain the answer to the question, you MUST state: "I'm sorry, but I cannot find that specific information in the report you've provided. It would be best to discuss this with your doctor."

---
**Context from uploaded documents:**
{context}
---
**User's Question:**
{question}
---
**Your Answer:**
"""                    
        PROMPT = PromptTemplate(template=prompt_template, input_variables=["context", "question"])

        ### CORRECTION 4: Updated chain creation to use the condense question prompt ###
        session["rag_chain"] = ConversationalRetrievalChain.from_llm(
            llm=llm,
            retriever=session["vector_store"].as_retriever(),
            condense_question_prompt=CONDENSE_QUESTION_PROMPT,
            combine_docs_chain_kwargs={"prompt": PROMPT},
            return_source_documents=False
        )

        session["uploaded_files"].append(file.filename)
        file_list = ", ".join(session['uploaded_files'])
        return {"status": "success", "message": f"Added '{file.filename}'. Knowledge base now contains: {file_list}"}

    except Exception as e:
        import traceback
        print(traceback.format_exc()) # Print full traceback to console for easier debugging
        raise HTTPException(status_code=500, detail=f"Failed to process file: {str(e)}")


@app.post("/chat", response_model=ChatResponse, summary="Handle a chat message for a specific session")
async def chat_with_bot(request: ChatRequest):
    if not request.session_id:
        raise HTTPException(status_code=400, detail="Session ID is missing.")
    
    session = sessions.get(request.session_id)

    # --- If RAG chain doesn't exist (no file uploaded), use the general medical persona ---
    if not session or not session.get("rag_chain"):
        llm = ChatGroq(temperature=0.7, model_name="llama-3.1-8b-instant", groq_api_key=os.getenv("GROQ_API_KEY"))
        
        medical_system_prompt = SystemMessage(
            content="""
            You are a specialized medical AI assistant named Dr. Asha. Your purpose is to provide general medical information.
            - When asked about your identity, state that you are a medical AI assistant.
            - You must politely decline to answer questions that are not related to medicine, health, or biology.
            - If you are asked a question before a document is uploaded, politely state: "I can help with that once you upload a medical report. Please use the upload button to provide a document."
            - Always conclude every response with the mandatory disclaimer: 'I am an AI assistant and not a substitute for professional medical advice. Please consult a healthcare professional for any health concerns.'
            """
        )

        try:
            response = llm.invoke([
                medical_system_prompt,
                HumanMessage(content=request.question)
            ])
            return ChatResponse(answer=response.content)
        except Exception as e:
            print(f"Error during general chat for session {request.session_id}: {e}")
            raise HTTPException(status_code=500, detail=f"An error occurred in general chat: {str(e)}")

    # --- Use the RAG chain if documents have been uploaded ---
    try:
        result = session["rag_chain"].invoke({
            "question": request.question,
            "chat_history": request.chat_history
        })
        return ChatResponse(answer=result["answer"])
    except Exception as e:
        print(f"Error during RAG chat for session {request.session_id}: {e}")
        raise HTTPException(status_code=500, detail=f"An error occurred during RAG chat: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    # Use "main:app" if you run this file directly as `python main.py`
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=True)
